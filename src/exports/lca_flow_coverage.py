"""``LcaFlowCoverageExporter`` — write ``to_review/lca_flow_coverage.xlsx``.

One row per unique AGB biosphere LCA flow keyed by ``name`` (case-folded).
This matches the "2,318 unique flows" universe counted on the freshly
imported AGB CSV — collapsing compartment and unit variants of the same
substance. For every flow we record:

* whether any AGB exchange of this flow is linked (``yes`` / ``partial`` /
  ``no``) and the dominant target DB it lands on,
* the target flow it is mapped to (name, categories, CAS),
* its presence in each sheet of ``placeholder_flow_classification.xlsx``
  (``ei`` / ``ef`` / ``neither``) — so reviewers can see novel AGB flows
  the placeholder workbook does not yet cover.

A second sheet lists placeholder rows whose flow ``name`` does not appear
in the AGB universe at all (stale / removed flows).

Columns mirror the EI-sheet layout of ``MappingsComparisonExporter``
(name / categories / unit / code / cas / matched_* / match_type /
category_proxy_match) with ``linked`` and ``target_database`` added so
the unlinked dimension has somewhere to live.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from config import Settings
from core.logging import Logging
from matching.bio_catalog import BiosphereCatalog
from readers import XlsxReader


@dataclass(frozen=True)
class _FlowRecord:
    """Aggregated view of every AGB biosphere exchange sharing one flow name."""

    name: str
    units: tuple[str, ...]
    categories: tuple[tuple[str, ...], ...]
    cas: str
    n_exchanges: int
    n_linked: int
    target_counts: tuple[tuple[tuple[str, str], int], ...]

    @property
    def linked_label(self) -> str:
        if self.n_linked == 0:
            return "no"
        if self.n_linked == self.n_exchanges:
            return "yes"
        return "partial"

    @property
    def dominant_target(self) -> tuple[str, str] | None:
        if not self.target_counts:
            return None
        # Highest count, ties broken by lex (db, code) for determinism.
        return sorted(self.target_counts, key=lambda kv: (-kv[1], kv[0][0], kv[0][1]))[0][0]

    @property
    def category_proxy_match(self) -> bool:
        """True when the source flow is observed in more than one compartment."""
        return len(self.categories) > 1


@dataclass(frozen=True)
class LcaFlowCoverageExporter:
    """Generate ``to_review/lca_flow_coverage.xlsx`` from ``sp.data``."""

    settings: Settings
    sp_data: list[dict]

    COLUMNS = (
        "name",
        "categories",
        "unit",
        "code",
        "cas",
        "linked",
        "target_database",
        "matched_name",
        "matched_categories",
        "matched_cas",
        "match_type",
        "category_proxy_match",
        "n_exchanges",
        "n_exchanges_linked",
        "in_placeholder_ei",
        "in_placeholder_ef",
        "in_placeholder_neither",
    )
    WIDTHS = (46, 42, 14, 18, 14, 10, 38, 46, 42, 14, 28, 14, 14, 16, 18, 18, 22)

    PLACEHOLDER_ONLY_COLUMNS = (
        "name",
        "categories",
        "unit",
        "code",
        "cas",
        "placeholder_sheet",
    )
    PLACEHOLDER_ONLY_WIDTHS = (46, 42, 14, 18, 14, 32)

    BANNER_FILL = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    LINKED_FILL = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    PARTIAL_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    UNLINKED_FILL = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
    PLACEHOLDER_ONLY_FILL = PatternFill(start_color="CFE2F3", end_color="CFE2F3", fill_type="solid")

    PLACEHOLDER_SHEETS = (
        ("ei", "match with ecoinvent v3.9.1"),
        ("ef", "match with EF v3.1"),
        ("neither", "Neither in ecoinvent nor EF"),
    )

    @property
    def _log(self):
        return Logging.get(__name__)

    @property
    def _ei_bio_db(self) -> str:
        return f"ecoinvent-{self.settings.resolved_ecoinvent_version}-biosphere"

    @property
    def _ef_db(self) -> str:
        return self.settings.ef_db_name

    # ------------------------------------------------------------------
    # Public entrypoint.

    def export(self) -> Path:
        s = self.settings
        out = s.paths.to_review / "lca_flow_coverage.xlsx"
        out.parent.mkdir(parents=True, exist_ok=True)

        flows = self._aggregate_flows()
        target_lookups = self._build_target_lookups()
        placeholder = self._load_placeholder_sheets()
        placeholder_names = self._placeholder_name_keys(placeholder)

        agb_rows = self._build_agb_rows(flows, target_lookups, placeholder_names)
        placeholder_only_rows = self._build_placeholder_only_rows(
            flows, placeholder, placeholder_names
        )

        wb = Workbook()
        wb.remove(wb.active)
        self._write_summary(wb, agb_rows, placeholder_only_rows, placeholder_names)
        self._write_agb_sheet(wb, agb_rows)
        self._write_placeholder_only_sheet(wb, placeholder_only_rows)
        wb.save(out)

        self._log.info(
            "lca_flow_coverage.written",
            path=str(out),
            n_flows=len(agb_rows),
            n_linked=sum(1 for r in agb_rows if r["linked"] == "yes"),
            n_partial=sum(1 for r in agb_rows if r["linked"] == "partial"),
            n_unlinked=sum(1 for r in agb_rows if r["linked"] == "no"),
            n_placeholder_only=len(placeholder_only_rows),
        )
        return out

    # ------------------------------------------------------------------
    # Aggregation: walk sp.data once, collapse by flow name.

    def _aggregate_flows(self) -> dict[str, _FlowRecord]:
        accum: dict[str, dict] = {}
        for proc in self.sp_data:
            for exc in proc.get("exchanges", []):
                if exc.get("type") != "biosphere":
                    continue
                name = str(exc.get("name") or "").strip()
                if not name:
                    continue
                key = name.lower()
                rec = accum.get(key)
                if rec is None:
                    rec = {
                        "name": name,
                        "units": set(),
                        "categories": set(),
                        "cas": "",
                        "n": 0,
                        "linked": 0,
                        "targets": Counter(),
                    }
                    accum[key] = rec
                rec["n"] += 1
                unit = str(exc.get("unit") or "").strip()
                if unit:
                    rec["units"].add(unit)
                cats = self._coerce_cats(exc.get("categories") or exc.get("context"))
                rec["categories"].add(cats)
                cas = self._norm_cas(exc.get("cas_number") or exc.get("CAS number"))
                if cas and not rec["cas"]:
                    rec["cas"] = cas
                inp = exc.get("input")
                if isinstance(inp, tuple) and len(inp) == 2:
                    rec["linked"] += 1
                    rec["targets"][inp] += 1
        return {k: self._finalise_flow(rec) for k, rec in accum.items()}

    @staticmethod
    def _finalise_flow(rec: dict) -> _FlowRecord:
        return _FlowRecord(
            name=rec["name"],
            units=tuple(sorted(rec["units"])),
            categories=tuple(sorted(rec["categories"])),
            cas=rec["cas"],
            n_exchanges=rec["n"],
            n_linked=rec["linked"],
            target_counts=tuple(rec["targets"].items()),
        )

    # ------------------------------------------------------------------
    # Target metadata: copied from MappingsComparisonExporter pattern.

    def _build_target_lookups(self) -> dict[str, dict[tuple[str, str], dict]]:
        catalog_path = self.settings.paths.registry_biosphere_catalog
        if not catalog_path.exists():
            return {}
        catalog = BiosphereCatalog.load(
            catalog_path,
            db_names=(self._ei_bio_db, self._ef_db, "biosphere3"),
        )
        out: dict[str, dict[tuple[str, str], dict]] = defaultdict(dict)
        for flow in catalog.flows:
            out[flow.db][(flow.db, flow.code)] = {
                "name": flow.name,
                "categories": tuple(flow.categories) if flow.categories else (),
                "cas": self._norm_cas(flow.cas),
                "unit": flow.unit,
            }
        return dict(out)

    # ------------------------------------------------------------------
    # Placeholder workbook.

    def _load_placeholder_sheets(self) -> dict[str, pd.DataFrame]:
        path = self.settings.paths.placeholder_xlsx
        if not path.exists():
            return {}
        xlsx = XlsxReader()
        return {
            label: xlsx.read(path, sheet_name=sheet) for label, sheet in self.PLACEHOLDER_SHEETS
        }

    @staticmethod
    def _placeholder_name_keys(sheets: dict[str, pd.DataFrame]) -> dict[str, set[str]]:
        out: dict[str, set[str]] = {}
        for label, df in sheets.items():
            if "name" not in df.columns:
                out[label] = set()
                continue
            out[label] = set(df["name"].dropna().astype(str).str.strip().str.lower())
        return out

    # ------------------------------------------------------------------
    # AGB-spine row builder.

    def _build_agb_rows(
        self,
        flows: dict[str, _FlowRecord],
        target_lookups: dict[str, dict[tuple[str, str], dict]],
        placeholder_names: dict[str, set[str]],
    ) -> list[dict]:
        rows: list[dict] = []
        for key, rec in sorted(flows.items()):
            target = rec.dominant_target
            if target:
                flow_meta = target_lookups.get(target[0], {}).get(target) or {}
                target_db = target[0]
                matched_name = flow_meta.get("name", "")
                matched_cats = flow_meta.get("categories", ())
                matched_cas = flow_meta.get("cas", "")
                match_type = self._classify_match_type(target_db)
            else:
                target_db = ""
                matched_name = ""
                matched_cats = ()
                matched_cas = ""
                match_type = "unlinked"
            rows.append(
                {
                    "name": rec.name,
                    "categories": self._render_categories(rec.categories),
                    "unit": ", ".join(rec.units),
                    "code": "",
                    "cas": rec.cas,
                    "linked": rec.linked_label,
                    "target_database": target_db,
                    "matched_name": matched_name,
                    "matched_categories": str(matched_cats) if matched_cats else "()",
                    "matched_cas": matched_cas,
                    "match_type": match_type,
                    "category_proxy_match": rec.category_proxy_match,
                    "n_exchanges": rec.n_exchanges,
                    "n_exchanges_linked": rec.n_linked,
                    "in_placeholder_ei": "yes" if key in placeholder_names.get("ei", set()) else "",
                    "in_placeholder_ef": "yes" if key in placeholder_names.get("ef", set()) else "",
                    "in_placeholder_neither": (
                        "yes" if key in placeholder_names.get("neither", set()) else ""
                    ),
                }
            )
        return rows

    def _classify_match_type(self, target_db: str) -> str:
        if target_db == self._ei_bio_db:
            return "linked_to_ecoinvent_biosphere"
        if target_db == self._ef_db:
            return "linked_to_ef"
        if target_db == "biosphere3":
            return "linked_to_biosphere3"
        return f"linked_to_{target_db}"

    # ------------------------------------------------------------------
    # Placeholder-only sheet (rows in placeholder whose flow name is not in AGB).

    def _build_placeholder_only_rows(
        self,
        flows: dict[str, _FlowRecord],
        placeholder: dict[str, pd.DataFrame],
        placeholder_names: dict[str, set[str]],
    ) -> list[dict]:
        agb_keys = set(flows.keys())
        rows: list[dict] = []
        for label, df in placeholder.items():
            if df.empty or "name" not in df.columns:
                continue
            for _, r in df.iterrows():
                name = str(r.get("name") or "").strip()
                if not name:
                    continue
                if name.lower() in agb_keys:
                    continue
                rows.append(
                    {
                        "name": name,
                        "categories": str(r.get("categories") or ""),
                        "unit": str(r.get("unit") or ""),
                        "code": str(r.get("code") or ""),
                        "cas": self._norm_cas(r.get("cas")),
                        "placeholder_sheet": label,
                    }
                )
        rows.sort(key=lambda r: (r["placeholder_sheet"], r["name"]))
        return rows

    # ------------------------------------------------------------------
    # Workbook writing.

    def _write_summary(
        self,
        wb: Workbook,
        agb_rows: list[dict],
        placeholder_only_rows: list[dict],
        placeholder_names: dict[str, set[str]],
    ) -> None:
        ws = wb.create_sheet("Summary")
        ws.append(["Metric", "Count"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = self.BANNER_FILL

        n_total = len(agb_rows)
        n_linked = sum(1 for r in agb_rows if r["linked"] == "yes")
        n_partial = sum(1 for r in agb_rows if r["linked"] == "partial")
        n_unlinked = sum(1 for r in agb_rows if r["linked"] == "no")
        in_ph_any = sum(
            1
            for r in agb_rows
            if any(r[f"in_placeholder_{k}"] == "yes" for k in ("ei", "ef", "neither"))
        )
        novel = n_total - in_ph_any

        rows: list[tuple[str, int]] = [
            ("AGB unique flow names (spine)", n_total),
            ("  fully linked", n_linked),
            ("  partially linked", n_partial),
            ("  fully unlinked", n_unlinked),
            ("AGB flows present in placeholder workbook (any sheet)", in_ph_any),
            (
                "  in `match with ecoinvent v3.9.1`",
                sum(1 for r in agb_rows if r["in_placeholder_ei"] == "yes"),
            ),
            (
                "  in `match with EF v3.1`",
                sum(1 for r in agb_rows if r["in_placeholder_ef"] == "yes"),
            ),
            (
                "  in `Neither in ecoinvent nor EF`",
                sum(1 for r in agb_rows if r["in_placeholder_neither"] == "yes"),
            ),
            ("AGB flows novel (not in any placeholder sheet)", novel),
            ("Placeholder rows whose name is absent from AGB", len(placeholder_only_rows)),
            ("  placeholder unique names: ei", len(placeholder_names.get("ei", set()))),
            ("  placeholder unique names: ef", len(placeholder_names.get("ef", set()))),
            ("  placeholder unique names: neither", len(placeholder_names.get("neither", set()))),
        ]
        for label, n in rows:
            ws.append([label, n])
        ws.column_dimensions["A"].width = 60
        ws.column_dimensions["B"].width = 12
        ws.freeze_panes = "A2"

    def _write_agb_sheet(self, wb: Workbook, rows: list[dict]) -> None:
        ws = wb.create_sheet("AGB flows")
        ws.append(list(self.COLUMNS))
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True)
            cell.fill = self.BANNER_FILL
        fills = {"yes": self.LINKED_FILL, "partial": self.PARTIAL_FILL, "no": self.UNLINKED_FILL}
        for row in rows:
            ws.append([_render(row.get(c, "")) for c in self.COLUMNS])
            fill = fills.get(row["linked"], self.UNLINKED_FILL)
            for cell in ws[ws.max_row]:
                cell.fill = fill
        for i, w in enumerate(self.WIDTHS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = "A2"

    def _write_placeholder_only_sheet(self, wb: Workbook, rows: list[dict]) -> None:
        ws = wb.create_sheet("Placeholder-only")
        ws.append(list(self.PLACEHOLDER_ONLY_COLUMNS))
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True)
            cell.fill = self.BANNER_FILL
        for row in rows:
            ws.append([_render(row.get(c, "")) for c in self.PLACEHOLDER_ONLY_COLUMNS])
            for cell in ws[ws.max_row]:
                cell.fill = self.PLACEHOLDER_ONLY_FILL
        for i, w in enumerate(self.PLACEHOLDER_ONLY_WIDTHS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # Pure helpers.

    @staticmethod
    def _coerce_cats(value) -> tuple[str, ...]:
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return ()
        if isinstance(value, (tuple, list)):
            return tuple(str(x) for x in value)
        return ()

    @staticmethod
    def _render_categories(cats: tuple[tuple[str, ...], ...]) -> str:
        if not cats:
            return "()"
        if len(cats) == 1:
            return str(cats[0]) if cats[0] else "()"
        return "; ".join(str(c) if c else "()" for c in cats)

    @staticmethod
    def _norm_cas(value) -> str:
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return ""
        s = str(value).strip()
        if not s:
            return ""
        digits = s.replace("-", "").lstrip("0") or "0"
        return digits


def _render(value) -> str:
    """Workbook-cell-safe coercion."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    if isinstance(value, (tuple, list)):
        return str(value)
    return str(value)
